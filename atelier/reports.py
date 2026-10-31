"""Human-readable Excel reports; all currency stored in minor units (kopecks)."""
from __future__ import annotations

from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .store import Store


def safe_cell(value: object) -> object:
    """Prevent user input from being interpreted as an Excel formula."""
    if isinstance(value,str) and value.lstrip(" \t\r\n\ufeff\u200b").startswith(("=","+","-","@")):
        return "'"+value
    return value


def make_report(store: Store) -> bytes:
    orders,payments,variants=store.export_rows()
    summary=store.analytics()
    book=Workbook()
    sh=book.active
    sh.title="Итоги"
    sh.append(["Показатель","Рубли","Пояснение"])
    for key,label,meaning in [
        ("net_receipts_cents","Получено за вычетом возвратов","Только вручную подтверждённые поступления и возвраты"),
        ("completed_revenue_cents","Стоимость завершённых заказов","Стоимость отгруженных/выданных и завершённых заказов"),
        ("completed_cost_cents","Себестоимость завершённых заказов","Учётная оценка из карточек товаров или сметы"),
        ("completed_gross_margin_cents","Валовая маржа завершённых заказов","Стоимость завершённых заказов минус их себестоимость; до накладных расходов и налогов"),
        ("active_orders_value_cents","Сумма неотменённых заказов","Обязательства, не фактическая выручка"),
    ]:
        sh.append([label,summary[key]/100,meaning])
        sh.cell(sh.max_row,2).number_format='#,##0.00 "₽"'
    sh.append(["Всего заказов",summary["order_count"],"Включая отменённые"])
    tab=book.create_sheet("Заказы")
    tab.append(["№", "Клиент Telegram ID", "Тип", "Статус", "Цена ₽", "Себестоимость ₽",
                "Оплачено ₽", "Остаток к оплате ₽", "Контакт", "Телефон",
                "Адрес/доставка", "Пожелания", "Мерки", "Создан"])
    for o in orders:
        tab.append([o["id"], o["telegram_id"], o["kind"], o["status"], o["total_cents"] / 100,
                    o["cost_cents"] / 100, o["paid_cents"] / 100,
                    (0 if o["status"] == "canceled" else o["total_cents"] - o["paid_cents"]) / 100,
                    safe_cell(o["contact_name"]), safe_cell(o["contact_phone"]),
                    safe_cell(o["delivery_address"]), safe_cell(o["custom_details"]),
                    safe_cell(o["custom_measurements"]), o["created_at"]])
        for col in (5, 6, 7, 8):
            tab.cell(tab.max_row, col).number_format = '#,##0.00 "₽"'
    tab=book.create_sheet("Платежи")
    tab.append(["№ платежа","№ заказа","Операция","Сумма ₽","Комментарий","Время","Идентификатор"])
    for p in payments:
        tab.append([p["id"],p["order_id"],p["kind"],p["amount_cents"]/100,
                    safe_cell(p["note"]),p["created_at"],safe_cell(p["reference"])])
        tab.cell(tab.max_row,4).number_format='#,##0.00 "₽"'
    tab=book.create_sheet("Ассортимент")
    tab.append(["SKU","Товар","Размер","Цвет","Цена ₽","Себестоимость ₽","Остаток","Активен"])
    for v in variants:
        tab.append([safe_cell(v["sku"]),safe_cell(v["product_name"]),safe_cell(v["size"]),safe_cell(v["color"]),
                    v["price_cents"]/100,v["cost_cents"]/100,v["stock"],bool(v["active"])])
        for col in (5,6): tab.cell(tab.max_row,col).number_format='#,##0.00 "₽"'
    for ws in book:
        ws.freeze_panes="A2"
        ws.auto_filter.ref=ws.dimensions
        ws.row_dimensions[1].height=28
        for cell in ws[1]:
            cell.font=Font(name="Calibri",bold=True,color="FFFFFF")
            cell.fill=PatternFill("solid",fgColor="203452")
            cell.alignment=Alignment(wrap_text=True,vertical="center")
        for col in ws.columns:
            letter=get_column_letter(col[0].column)
            width=min(50,max(13,max(len(str(c.value or "")) for c in list(col)[:80])+2))
            ws.column_dimensions[letter].width=width
    data=BytesIO()
    book.save(data)
    return data.getvalue()
