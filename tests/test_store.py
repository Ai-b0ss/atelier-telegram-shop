import sqlite3
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO

import pytest
from openpyxl import load_workbook

from atelier.reports import make_report, safe_cell
from atelier.store import Store, StoreError


@pytest.fixture
def shop(tmp_path):
    store = Store(tmp_path / "data" / "shop.sqlite3")
    store.seed_demo()
    return store


def client(shop, tg=101):
    return shop.customer(tg, f"Customer {tg}", "client")


def first_variant(shop):
    with shop.connection() as conn:
        return conn.execute("SELECT id FROM variants ORDER BY id LIMIT 1").fetchone()[0]


def test_seed_idempotent(shop):
    shop.seed_demo()
    assert len(shop.categories()) == 3
    assert sum(row["items"] for row in shop.categories()) == 3


def test_cart_checkout_snapshot_and_stock(shop):
    cid = client(shop)
    vid = first_variant(shop)
    shop.add_to_cart(cid, vid, 2)
    shop.add_to_cart(cid, vid)
    assert shop.cart(cid)[0]["quantity"] == 3
    order = shop.checkout(cid, "Ольга", "+79990001122", "Самовывоз")
    assert shop.cart(cid) == []
    with shop.connection() as conn:
        variant = conn.execute("SELECT * FROM variants WHERE id=?", (vid,)).fetchone()
        item = conn.execute("SELECT * FROM order_items WHERE order_id=?", (order,)).fetchone()
        assert variant["stock"] == 5
        assert item["quantity"] == 3
        assert item["unit_price_cents"] == variant["price_cents"]
    assert shop.order(order)["total_cents"] == 3 * item["unit_price_cents"]


def test_empty_cart_and_invalid_contacts(shop):
    cid = client(shop)
    with pytest.raises(StoreError, match="пуста"):
        shop.checkout(cid, "Иван", "+7", "самовывоз")
    shop.add_to_cart(cid, first_variant(shop))
    with pytest.raises(StoreError, match="Укажите"):
        shop.checkout(cid, "", "+7", "самовывоз")
    assert len(shop.cart(cid)) == 1


def test_stock_race_across_customers(shop):
    vid = first_variant(shop)
    with shop.connection() as conn:
        conn.execute("UPDATE variants SET stock=1 WHERE id=?", (vid,))
    a, b = client(shop, 101), client(shop, 202)
    shop.add_to_cart(a, vid)
    shop.add_to_cart(b, vid)
    def buy(cid):
        try:
            return shop.checkout(cid, "Заказчик", "@contact", "доставка")
        except StoreError:
            return None
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(buy, [a, b]))
    assert len([r for r in results if r is not None]) == 1
    with shop.connection() as conn:
        assert conn.execute("SELECT stock FROM variants WHERE id=?", (vid,)).fetchone()[0] == 0
        assert conn.execute("SELECT COUNT(*) FROM orders").fetchone()[0] == 1


def test_stock_cancel_restores_once(shop):
    cid=client(shop);vid=first_variant(shop)
    with shop.connection() as conn:
        original=conn.execute("SELECT stock FROM variants WHERE id=?",(vid,)).fetchone()[0]
    shop.add_to_cart(cid,vid)
    oid=shop.checkout(cid,"Даша","@dasha","Самовывоз")
    shop.set_status(oid,"canceled")
    with pytest.raises(StoreError): shop.set_status(oid,"canceled")
    with shop.connection() as conn:
        assert conn.execute("SELECT stock FROM variants WHERE id=?",(vid,)).fetchone()[0]==original


def test_custom_quote_payment_refund_and_cancel(shop):
    cid=client(shop)
    oid=shop.custom_order(cid,"Мария","@maria","Москва","Платье с вышивкой по эскизу","Талия 70 см")
    with pytest.raises(StoreError,match="цену"):
        shop.set_status(oid,"confirmed")
    shop.set_custom_price(oid,1200000,500000)
    shop.set_status(oid,"confirmed")
    shop.record_payment(oid,"received",400000)
    with pytest.raises(StoreError,match="полученную"):
        shop.record_payment(oid,"refunded",500000)
    with pytest.raises(StoreError,match="возврат"):
        shop.set_status(oid,"canceled")
    shop.record_payment(oid,"refunded",400000)
    shop.set_status(oid,"canceled")
    assert shop.order(oid)["status"]=="canceled"
    assert shop.analytics()["net_receipts_cents"]==0


def test_full_lifecycle_and_financial_report(shop):
    cid=client(shop)
    shop.add_to_cart(cid,first_variant(shop))
    oid=shop.checkout(cid,"Мария","@maria","Мск")
    amount=shop.order(oid)["total_cents"]
    cost=shop.order(oid)["cost_cents"]
    shop.record_payment(oid,"received",amount)
    with pytest.raises(StoreError): shop.record_payment(oid,"received",1)
    for status in ("confirmed","in_production","ready","shipped","completed"):
        shop.set_status(oid,status)
    with pytest.raises(StoreError): shop.set_status(oid,"canceled")
    a=shop.analytics()
    assert a["net_receipts_cents"]==amount
    assert a["completed_cost_cents"]==cost
    assert a["indicative_margin_cents"]==amount-cost
    xlsx=make_report(shop)
    book=load_workbook(BytesIO(xlsx))
    assert book.sheetnames==["Итоги","Заказы","Платежи","Ассортимент"]
    assert book["Итоги"]["B2"].value==amount/100
    assert book["Платежи"]["D2"].value==amount/100
    assert book["Заказы"]["D2"].value=="completed"


def test_csv_excel_formula_sanitization(shop):
    assert safe_cell("=1+1")=="'=1+1"
    assert safe_cell("+cmd")=="'+cmd"
    cid=client(shop)
    oid=shop.custom_order(cid,"=HYPERLINK(\"x\")","@a","@somewhere","Пальто с воротником","Сантиметры")
    book=load_workbook(BytesIO(make_report(shop)))
    cell=book["Заказы"]["G2"]
    assert cell.data_type=="s" and cell.value.startswith("'=")


def test_custom_quote_constraints_and_status_filter(shop):
    cid=client(shop)
    oid=shop.custom_order(cid,"Вика","@v","Самовывоз","Блузка с интересным воротником","не знаю")
    with pytest.raises(StoreError): shop.set_custom_price(oid,0,10)
    shop.set_custom_price(oid,10000,11000)
    assert [r["id"] for r in shop.list_orders("new")]==[oid]
    assert shop.list_orders("completed")==[]
    with pytest.raises(StoreError): shop.list_orders("UNKNOWN")
    with pytest.raises(StoreError): shop.record_payment(oid,"invalid",100)


def test_order_snapshot_price_after_variant_change(shop):
    cid=client(shop);vid=first_variant(shop)
    shop.add_to_cart(cid,vid)
    oid=shop.checkout(cid,"Лена","@lena","Мск")
    price=shop.order(oid)["total_cents"]
    with shop.connection() as db:
        db.execute("UPDATE variants SET price_cents=price_cents*2 WHERE id=?",(vid,))
    assert shop.order(oid)["total_cents"]==price


def test_foreign_keys_and_nonnegative_stock(shop):
    with pytest.raises(sqlite3.IntegrityError):
        with shop.connection() as db:
            db.execute("INSERT INTO cart_items(customer_id,variant_id,quantity) VALUES(99999,99999,1)")
    with pytest.raises(sqlite3.IntegrityError):
        with shop.connection() as db:
            db.execute("UPDATE variants SET stock=-1 WHERE id=?",(first_variant(shop),))


def test_amount_parser_from_catalogue_cli():
    from atelier.manage import money_to_cents
    assert money_to_cents("1000.25") == 100025
    assert money_to_cents("10,50") == 1050
    for bad in ("nan", "Infinity", "-3", "1.001", "xx"):
        with pytest.raises(StoreError): money_to_cents(bad)


def test_status_history_has_all_transitions(shop):
    cid = client(shop)
    shop.add_to_cart(cid, first_variant(shop))
    oid = shop.checkout(cid, "Анна", "@anna", "самовывоз")
    shop.set_status(oid, "confirmed")
    shop.set_status(oid, "ready")
    with shop.connection() as db:
        states = db.execute("SELECT old_status,new_status FROM status_history WHERE order_id=? ORDER BY id", (oid,)).fetchall()
    assert [(s[0],s[1]) for s in states] == [(None,"new"),("new","confirmed"),("confirmed","ready")]
